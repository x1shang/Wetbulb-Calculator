"""Non-interactive GUI smoke check, also available in the packaged application."""
import json
from pathlib import Path
import tempfile


def run(window, app, about_class, unit_class, output):
    import core
    from batch import calculate_file
    from openpyxl import Workbook, load_workbook

    errors = []
    window.createErrorInfoBar = lambda message, *a, **k: errors.append(message)
    window.show()
    assert not window.windowIcon().isNull(), 'Missing application icon'
    counts = []
    for mode, value in enumerate(('15', '20', '60')):
        window.ComboBox.setCurrentIndex(mode)
        window.LineEdit_3.setText('25')
        window.LineEdit.setText(value)
        window.LineEdit_2.setText('1013.25')
        window.validate_and_calculate()
        app.processEvents()
        rows = window.list_model.stringList()
        assert len(rows) == 16, (mode, rows)
        assert any(r.startswith('Goff-水面') and '℃' in r for r in rows)
        counts.append(len(rows))
    about = about_class()
    about.show()
    unit = unit_class(window)
    unit.show()
    app.processEvents()
    assert core.tag in about.label.text()
    assert about.validation.text() == '经过IAPWS-95 + MK2005 + Stull + ASHRAE + Bolton五源交叉验证'
    with tempfile.TemporaryDirectory(prefix='wetbulb-smoke-') as directory:
        source, target = Path(directory) / 'input.xlsx', Path(directory) / 'output.xlsx'
        book = Workbook()
        book.active.append(['A', 'B', 'C'])
        book.active.append([25, 60, 1013.25])
        book.active.append([25, -1, 1013.25])
        book.save(source)
        book.close()
        calculate_file(source, target, mode=2)
        book = load_workbook(target)
        assert isinstance(book.active['D2'].value, float)
        assert book.active['D3'].data_type == 'e'
        book.close()
    assert not errors, errors
    Path(output).write_text(json.dumps({'version': core.tag, 'rows': counts, 'batch': 'ok'}), encoding='utf-8')
    unit.close()
    about.close()
    window.close()
