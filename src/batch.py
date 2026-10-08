"""Spreadsheet calculation shared by the GUI and end-to-end tests."""
from pathlib import Path

from openpyxl import load_workbook

import core


def calculate_file(source, output, mode=0, to_celsius=float, to_hpa=float,
                   from_celsius=float, temperature_unit='℃', pressure_unit='hPa',
                   progress=None):
    if mode not in (0, 1, 2):
        raise ValueError('未知计算模式')
    if Path(source).resolve() == Path(output).resolve():
        raise ValueError('输出文件不能覆盖输入文件')
    book = load_workbook(source)
    try:
        sheet = book.active
        headers = [cell.value for cell in sheet[1]]
        if not all(name in headers for name in ('A', 'B', 'C')):
            raise ValueError('Excel文件必须包含ABC列！')
        cols = [headers.index(name) + 1 for name in ('A', 'B', 'C')]
        result_col = max(sheet.max_column + 1, 4)
        names = ('湿球',) if mode == 0 else ('露点',) if mode == 1 else ('露点', '湿球')
        for offset, name in enumerate(names):
            sheet.cell(1, result_col + offset, name + temperature_unit)
        error_col = result_col + len(names)
        sheet.cell(1, error_col, '错误')
        for row in range(2, sheet.max_row + 1):
            try:
                a, b, c = [sheet.cell(row, col).value for col in cols]
                t, p = to_celsius(float(a)), to_hpa(float(c))
                value = float(b) if mode == 2 else to_celsius(float(b))
                if mode == 0:
                    results = core.calculate_wetbulb(value, t, value, p)
                elif mode == 1:
                    results = core.calculate_dewpoint(t, value, p)
                else:
                    results = core.calculate_both(t, t, value, p)
                method = 'Goff-水面' if t >= 0 else 'Goff-冰面'
                result = next(r for r in results if r['method'] == method)
                values = [result['result1']] if mode != 2 else [result['result1'], result['result2']]
                if not all(isinstance(v, (float, int)) for v in values):
                    raise ValueError(' / '.join(str(v) for v in values))
                for offset, value in enumerate(values):
                    sheet.cell(row, result_col + offset, from_celsius(value))
            except (ValueError, TypeError, OverflowError) as exc:
                for offset in range(len(names)):
                    cell = sheet.cell(row, result_col + offset, '#VALUE!')
                    cell.data_type = 'e'
                # Force text: an input/error message must never become an Excel formula.
                cell = sheet.cell(row, error_col, str(exc))
                cell.data_type = 's'
            if progress:
                progress(row - 1, sheet.max_row - 1)
        labels = ('干球' + temperature_unit,
                  ('露点' + temperature_unit, '湿球' + temperature_unit, '相对湿度%')[mode],
                  '大气' + pressure_unit)
        for col, label in zip(cols, labels):
            sheet.cell(1, col, label)
        book.save(output)
    finally:
        book.close()
    return Path(output)
