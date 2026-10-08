"""Exercise the real CLI process and GUI spreadsheet entry point."""
import json
import os
from pathlib import Path
import subprocess
import sys

import pytest
import core
from test_repo_hygiene import _gui_interpreter, _run_probe, _probe_failed

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.parametrize('mode,value', [('dewpoint', '15'), ('wetbulb', '20'), ('rh', '60')])
def test_cli_without_site_packages(mode, value):
    proc = subprocess.run([sys.executable, '-S', 'main.py', '--cli', '--mode', mode,
                           '--temperature', '25', '--value', value], cwd=ROOT,
                          capture_output=True, encoding='utf-8',
                          env={**os.environ, 'PYTHONIOENCODING': 'utf-8'}, timeout=30)
    assert proc.returncode == 0, proc.stderr
    payload = json.loads(proc.stdout)
    assert payload['version'] == core.tag
    assert len(payload['results']) == 14
    goff = next(r for r in payload['results'] if r['method'] == 'Goff-水面')
    assert isinstance(goff['result1'], float)


@pytest.mark.parametrize('value', ['-1', '0', '101', 'nan', 'inf'])
def test_cli_rejects_invalid_humidity(value):
    proc = subprocess.run([sys.executable, '-S', 'src/cli.py', '--mode', 'rh',
                           '--temperature', '25', '--value', value], cwd=ROOT,
                          capture_output=True, encoding='utf-8',
                          env={**os.environ, 'PYTHONIOENCODING': 'utf-8'}, timeout=30)
    assert proc.returncode == 2
    assert json.loads(proc.stderr)['error']
    assert not proc.stdout


@pytest.mark.parametrize('mode,value', [(0, 15), (1, 20), (2, 60)])
def test_gui_batch_writes_valid_and_error_cells(mode, value):
    exe = _gui_interpreter()
    if exe is None:
        pytest.skip('GUI environment unavailable')
    code = f'''
import os, tempfile
os.environ['QT_QPA_PLATFORM'] = 'offscreen'
import main
import core
from pathlib import Path
from openpyxl import Workbook, load_workbook
from PySide2.QtWidgets import QApplication
app = QApplication([])
w = main.main_window()
w.ComboBox.setCurrentIndex({mode})
errors, successes = [], []
w.createErrorInfoBar = lambda message, *a, **k: errors.append(message)
w.createSuccessInfoBar = lambda message, *a, **k: successes.append(message)
with tempfile.TemporaryDirectory() as td:
    main._HERE = td
    source = Path(td) / 'sample.xlsx'
    book = Workbook()
    book.active.append(['A', 'B', 'C'])
    book.active.append([25, {value}, 1013.25])
    book.active.append([25, 'invalid', 1013.25])
    book.active.append([25, {value}, 1013.25])
    book.save(source)
    book.close()
    w.process_excel_file()
    assert not errors, errors
    assert successes
    result = load_workbook(Path(td) / 'result_sample.xlsx')
    sheet = result.active
    assert sheet.max_row == 4
    expected = (core.calculate_wetbulb({value}, 25, {value}) if {mode} == 0 else
                core.calculate_dewpoint(25, {value}, 1013.25) if {mode} == 1 else
                core.calculate_both(25, 25, {value}))
    goff = next(r for r in expected if r['method'] == 'Goff-水面')
    assert abs(sheet['D2'].value - goff['result1']) < 1e-9
    assert sheet['D2'].value == sheet['D4'].value
    assert sheet['D3'].value == '#VALUE!' and sheet['D3'].data_type == 'e'
    if {mode} == 2:
        assert abs(sheet['E2'].value - goff['result2']) < 1e-9
        assert sheet['E3'].data_type == 'e'
    assert sheet.cell(3, 6 if {mode} == 2 else 5).value
    result.close()
    w.process_excel_file()  # Existing result_ files must not be mistaken for input.
    assert not errors, errors
    assert len(successes) == 2
    assert not w.ProgressBar.isVisible()
    assert w.pushButton_7.isEnabled()
print('BATCH_OK')
'''
    proc, records = _run_probe(exe, code)
    assert records, _probe_failed(proc, records)
    assert proc.returncode == 0, proc.stderr
    assert 'BATCH_OK' in records
